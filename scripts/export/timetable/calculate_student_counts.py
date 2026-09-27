#!/usr/bin/env python3
# Copyright (c) 2026 Filip Marić. See LICENCE.
"""Decode student counts for rows in a timetable ``people`` section."""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parents[2]
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))
BACKEND_DIR = SCRIPT_DIR.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))
from lib.timetable_common import CYRILLIC_TO_LATIN, normalize_group_name


OUTPUT_FIELDS = (
    "group_codes",
    "teacher_username",
    "course_alias",
    "course_type",
    "student_count",
)
COURSE_CODE_COLUMN = 2  # B
STUDENT_BLOCK_START_COLUMN = 21  # U


@dataclass(frozen=True)
class GroupInfo:
    study_year: str
    modules: frozenset[str]
    accreditation: str
    fraction: float


@dataclass(frozen=True)
class PeopleRow:
    group_codes: tuple[str, ...]
    teacher_username: str
    course_alias: str
    course_type: str


def normalize_code(value: object) -> str:
    text = str(value or "").translate(CYRILLIC_TO_LATIN).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", text)


def load_group_mapping(path: Path) -> dict[str, GroupInfo]:
    result: dict[str, GroupInfo] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            code = normalize_group_name(row.get("group_code"))
            if not code:
                continue
            modules = frozenset(
                module.strip().casefold()
                for module in str(row.get("module") or "").split(",")
                if module.strip()
            )
            result[code] = GroupInfo(
                study_year=str(row.get("study_year") or "").strip(),
                modules=modules,
                accreditation=str(row.get("accreditation") or "").strip(),
                fraction=float(str(row.get("fraction") or "1").replace(",", ".")),
            )
    return result


def read_people_rows(path: Path) -> list[PeopleRow]:
    rows: list[PeopleRow] = []
    in_people = False
    with path.open(encoding="utf-8-sig") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line:
                continue
            if line == "people":
                in_people = True
                continue
            if not in_people:
                continue
            if line.startswith("//"):
                continue
            if "\t" not in raw_line:
                break
            fields = raw_line.rstrip("\r\n").split("\t")
            if len(fields) < 3:
                continue
            groups = tuple(
                normalize_group_name(value)
                for value in re.split(r"[,.;|]+", fields[0])
                if normalize_group_name(value)
            )
            if groups:
                course_parts = fields[2].strip().split(".", 1)
                rows.append(
                    PeopleRow(
                        group_codes=groups,
                        teacher_username=fields[1].strip(),
                        course_alias=course_parts[0].strip().casefold(),
                        course_type=course_parts[1].strip().casefold() if len(course_parts) > 1 else "",
                    )
                )
    return rows


def parse_block(value: object) -> dict[str, str]:
    result: dict[str, str] = {}
    for raw_line in str(value or "").splitlines():
        if ":" not in raw_line:
            continue
        key, block_value = raw_line.split(":", 1)
        result[key.strip().casefold()] = block_value.strip()
    return result


def _number(value: str | None) -> float | None:
    if value is None or value.strip() in {"", "???", "-"}:
        return None
    normalized = value.strip().replace(" ", "").replace(",", ".")
    try:
        return float(normalized)
    except ValueError:
        return None


def load_course_rows(path: Path, aliases_path: Path) -> dict[str, list[tuple[str, int, list[dict[str, str]]]]]:
    aliases: dict[str, str] = {}
    with aliases_path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            original = normalize_code(row.get("original_code"))
            alias = str(row.get("timetable_code") or "").strip().casefold()
            if original and alias:
                aliases[alias] = original

    result: dict[str, list[tuple[str, int, list[dict[str, str]]]]] = {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in ("predmeti1sem", "predmeti2sem"):
        worksheet = workbook[sheet]
        for row_number, values in enumerate(worksheet.iter_rows(values_only=True), 1):
            if len(values) < STUDENT_BLOCK_START_COLUMN:
                continue
            original = normalize_code(values[COURSE_CODE_COLUMN - 1])
            if not original:
                continue
            blocks = [
                parse_block(value)
                for value in values[STUDENT_BLOCK_START_COLUMN - 1 :]
                if value is not None
            ]
            result.setdefault(original, []).append((sheet, row_number, blocks))
    return {alias: rows for alias, original in aliases.items() for rows in [result.get(original, [])]}


def matching_student_count(blocks: list[dict[str, str]], info: GroupInfo) -> float:
    total = 0.0
    wanted_accreditation = info.accreditation.casefold()
    wanted_year = info.study_year.casefold()
    for block in blocks:
        accreditation = block.get("акредитација", "").strip().casefold()
        year = block.get("година", "").strip().casefold()
        module = block.get("модул", "").strip().casefold()
        if accreditation != wanted_accreditation or year != wanted_year or module not in info.modules:
            continue
        count = _number(block.get("студената"))
        if count is not None:
            total += count
    return total


def decode(people_path: Path, mapping_path: Path, workbook_path: Path, aliases_path: Path) -> list[dict[str, str]]:
    mapping = load_group_mapping(mapping_path)
    course_rows = load_course_rows(workbook_path, aliases_path)
    decoded: list[dict[str, str]] = []
    for people_row in read_people_rows(people_path):
        source_rows = course_rows.get(people_row.course_alias, [])
        total = 0.0
        for group_code in people_row.group_codes:
            info = mapping.get(group_code)
            if info is None:
                continue
            for _, _, blocks in source_rows:
                total += matching_student_count(blocks, info) * info.fraction
        # Student counts are non-negative; use the usual half-up rule so 65.5 -> 66.
        student_count = str(int(total + 0.5))
        decoded.append(
            {
                "group_codes": ",".join(people_row.group_codes),
                "teacher_username": people_row.teacher_username,
                "course_alias": people_row.course_alias,
                "course_type": people_row.course_type,
                "student_count": student_count,
            }
        )
    return decoded


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("people_file", type=Path)
    parser.add_argument("mapping_csv", type=Path)
    parser.add_argument("courses_xlsx", type=Path)
    parser.add_argument("aliases_csv", type=Path)
    parser.add_argument("output_csv", type=Path)
    args = parser.parse_args()
    rows = decode(args.people_file, args.mapping_csv, args.courses_xlsx, args.aliases_csv)
    args.output_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Generated decoded people counts: {len(rows)} rows -> {args.output_csv}")


if __name__ == "__main__":
    main()
