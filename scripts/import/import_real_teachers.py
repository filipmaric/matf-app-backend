#!/usr/bin/env python3
# Copyright (c) 2026 Filip Marić. See LICENCE.

"""Auxiliary fallback importer for teacher names found in course_sessions.xlsx.

Use this instead of import_teachers.py only when the authoritative teachers.csv
source is unavailable. It is intended for one-off or backup database recovery.
"""

from __future__ import annotations

import argparse
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


EXCLUDED_SHEETS = {"Фонд", "Све"}
ROLE_MARKERS = ("Сарадник", "Наставник", "Асистент")


CYRILLIC_TO_LATIN = str.maketrans(
    {
        "А": "A",
        "Б": "B",
        "В": "V",
        "Г": "G",
        "Д": "D",
        "Ђ": "Dj",
        "Е": "E",
        "Ж": "Z",
        "З": "Z",
        "И": "I",
        "Ј": "J",
        "К": "K",
        "Л": "L",
        "Љ": "Lj",
        "М": "M",
        "Н": "N",
        "Њ": "Nj",
        "О": "O",
        "П": "P",
        "Р": "R",
        "С": "S",
        "Т": "T",
        "Ћ": "C",
        "У": "U",
        "Ф": "F",
        "Х": "H",
        "Ц": "C",
        "Ч": "C",
        "Џ": "Dz",
        "Ш": "S",
        "а": "a",
        "б": "b",
        "в": "v",
        "г": "g",
        "д": "d",
        "ђ": "dj",
        "е": "e",
        "ж": "z",
        "з": "z",
        "и": "i",
        "ј": "j",
        "к": "k",
        "л": "l",
        "љ": "lj",
        "м": "m",
        "н": "n",
        "њ": "nj",
        "о": "o",
        "п": "p",
        "р": "r",
        "с": "s",
        "т": "t",
        "ћ": "c",
        "у": "u",
        "ф": "f",
        "х": "h",
        "ц": "c",
        "ч": "c",
        "џ": "dz",
        "ш": "s",
    }
)


def transliterate(text: str) -> str:
    text = text.translate(CYRILLIC_TO_LATIN)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z0-9]+", "", text)
    return text.lower()


def derive_base_username(full_name: str) -> str:
    parts = [part for part in re.split(r"\s+", full_name.strip()) if part]
    if not parts:
        raise ValueError("empty teacher name")
    if len(parts) == 1:
        return transliterate(parts[0])
    return f"{transliterate(parts[0])}.{transliterate(parts[1])}"


def load_real_teacher_names(workbook_path: Path) -> list[str]:
    wb = load_workbook(workbook_path, data_only=True)
    names: set[str] = set()
    for sheet_name in wb.sheetnames:
        if sheet_name in EXCLUDED_SHEETS:
            continue
        ws = wb[sheet_name]
        for row in range(2, ws.max_row + 1):
            value = ws.cell(row=row, column=1).value
            if value in (None, ""):
                continue
            name = str(value).strip()
            if not name:
                continue
            compact = name.strip()
            has_role = any(marker in compact for marker in ROLE_MARKERS)
            has_parens = "(" in compact and ")" in compact
            is_all_caps = compact.isupper() and len(compact) > 1
            looks_like_unit_label = bool(re.match(r"^[А-ЯA-Z]{2,}\b", compact))
            looks_like_person = bool(re.search(r"[А-Я][а-я]+\s+[А-Я][а-я]+", compact))
            if has_role and has_parens:
                continue
            if has_role or is_all_caps or compact == "ФИЗИКА":
                continue
            if looks_like_unit_label and not looks_like_person:
                continue
            names.add(name)
    return sorted(names)


def ensure_schema(conn: sqlite3.Connection, schema_path: Path) -> None:
    has_teachers = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'teachers'"
    ).fetchone()
    if has_teachers:
        return
    if not schema_path.exists():
        raise FileNotFoundError(f"Schema file not found: {schema_path}")
    conn.executescript(schema_path.read_text(encoding="utf-8"))


def existing_usernames(conn: sqlite3.Connection) -> set[str]:
    return {
        str(row[0]).strip()
        for row in conn.execute("SELECT username FROM teachers WHERE username IS NOT NULL")
        if row[0] not in (None, "")
    }


def plan_import(
    names: list[str],
    used_usernames: set[str],
) -> tuple[list[tuple[str, str]], list[tuple[str, str, str]]]:
    planned: list[tuple[str, str]] = []
    skipped: list[tuple[str, str, str]] = []

    for name in names:
        username = derive_base_username(name)
        if username in used_usernames:
            skipped.append((name, username, "duplicate username"))
            continue
        used_usernames.add(username)
        planned.append((name, username))

    return planned, skipped


def import_teachers(conn: sqlite3.Connection, planned: list[tuple[str, str]]) -> int:
    inserted = 0
    with conn:
        cur = conn.cursor()
        for name, username in planned:
            cur.execute(
                """
                INSERT INTO teachers (name, username)
                VALUES (?, ?)
                """,
                (name, username),
            )
            inserted += 1
    return inserted


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="Import real teacher names from course_sessions.xlsx into a SQLite backup db."
    )
    parser.add_argument(
        "--workbook",
        required=True,
        help="Path to course_sessions.xlsx",
    )
    parser.add_argument(
        "--database",
        required=True,
        help="SQLite database path",
    )
    parser.add_argument(
        "--schema",
        required=True,
        help="Path to backend/schema.sql used when the database is empty",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print the derived usernames without writing to the database",
    )
    args = parser.parse_args(argv)

    workbook_path = Path(args.workbook)
    database_path = Path(args.database)
    schema_path = Path(args.schema)

    if not workbook_path.exists():
        print(f"Workbook not found: {workbook_path}", file=sys.stderr)
        return 2

    names = load_real_teacher_names(workbook_path)
    if not names:
        print("No real teacher names found.", file=sys.stderr)
        return 1

    used = set()
    if database_path.exists():
        conn = sqlite3.connect(database_path)
        try:
            table_exists = conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'teachers'"
            ).fetchone()
            if table_exists:
                used.update(existing_usernames(conn))
        finally:
            conn.close()

    planned, skipped = plan_import(names, used)

    if args.dry_run:
        for name, username in planned:
            print(f"{username}\t{name}")
        for name, username, reason in skipped:
            print(f"SKIP\t{username}\t{name}\t{reason}")
        print(f"COUNT={len(planned)}")
        print(f"SKIPPED={len(skipped)}")
        return 0

    database_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(database_path)
    try:
        ensure_schema(conn, schema_path)
        inserted = import_teachers(conn, planned)
        print(f"Imported {inserted} teachers into {database_path}")
        if skipped:
            print(f"Skipped {len(skipped)} teachers:")
            for name, username, reason in skipped:
                print(f"- {username}\t{name}\t{reason}")
        return 0
    except Exception as exc:
        conn.rollback()
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        conn.close()


if __name__ == "__main__":
    raise SystemExit(main())
