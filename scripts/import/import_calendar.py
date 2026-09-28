#!/usr/bin/env python3
# Copyright (c) 2026 Filip Marić. See LICENCE.

"""Import the academic calendar and its day kinds from an Excel workbook."""

from __future__ import annotations

import argparse
import re
import sqlite3
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

BACKEND_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BACKEND_DIR))

from calendar_common import validate_day_kind
from db import ensure_calendar_revision_schema, ensure_calendar_schema


DAY_NAMES = {
    "понедељак": 0,
    "уторак": 1,
    "среда": 2,
    "сриједа": 2,
    "четвртак": 3,
    "петак": 4,
    "субота": 5,
    "недеља": 6,
}

BLUE_COLORS = {"FF4A86E8", "FF4285F4"}
FILL_TO_KIND = {
    "FFFFFF00": "teaching",
    "FFFF0000": "non_working",
    "FFFFF2CC": "colloquium",
    "FFFBBC04": "makeup",
    "FFFF9900": "makeup",
}


def _rgb(cell) -> str | None:
    color = cell.fill.fgColor if cell.fill else None
    return color.rgb if color and color.type == "rgb" else None


def _workbook_start_year(workbook, override: int | None) -> int:
    if override is not None:
        return override
    title = str(workbook.active["A1"].value or "")
    match = re.search(r"(20\d{2})\s*/\s*\d{2}", title)
    if not match:
        raise ValueError("could not determine academic year from workbook title; use --academic-year-start")
    return int(match.group(1))


def _comment_day(cell) -> int:
    if not cell.comment:
        raise ValueError(f"makeup date {cell.coordinate} is missing its replacement weekday comment")
    text = cell.comment.text.lower()
    for name, weekday in DAY_NAMES.items():
        if name in text:
            return weekday
    raise ValueError(f"makeup date {cell.coordinate} has no recognized weekday comment")


def _kind_for_cell(cell) -> str:
    color = _rgb(cell)
    if color in BLUE_COLORS:
        return "exam"
    return FILL_TO_KIND.get(color, "non_working")


def read_calendar(workbook_path: Path, academic_year_start: int | None = None):
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    start_year = _workbook_start_year(workbook, academic_year_start)
    first_date = date(start_year, 10, 1)
    last_date = date(start_year + 1, 9, 30)
    rows_by_date = {
        (first_date + timedelta(days=offset)).isoformat(): ("non_working", -1)
        for offset in range((last_date - first_date).days + 1)
    }

    # The workbook lays out weeks in rows, with 1 October in F3.
    sheet = workbook.active
    for row in range(3, sheet.max_row + 1):
        for column in range(3, 10):
            cell = sheet.cell(row, column)
            if not isinstance(cell.value, (int, float)):
                continue
            current_date = first_date + timedelta(days=(row - 3) * 7 + (column - 6))
            if not (first_date <= current_date <= last_date):
                continue
            kind = validate_day_kind(_kind_for_cell(cell))
            week_day = _comment_day(cell) if kind == "makeup" else -1
            rows_by_date[current_date.isoformat()] = (kind, week_day)

    expected = (last_date - first_date).days + 1
    rows = [(date_value, kind, week_day) for date_value, (kind, week_day) in rows_by_date.items()]
    if len(rows) != expected:
        raise ValueError(f"workbook contains {len(rows)} calendar dates; expected {expected}")
    return rows


def import_calendar(database_path: Path, workbook_path: Path, schema_path: Path, academic_year_start: int | None = None):
    if not workbook_path.exists():
        raise FileNotFoundError(f"workbook not found: {workbook_path}")

    rows = read_calendar(workbook_path, academic_year_start)
    conn = sqlite3.connect(database_path)
    try:
        with conn:
            if not conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='days'").fetchone():
                if not schema_path.exists():
                    raise FileNotFoundError(f"schema file not found: {schema_path}")
                conn.executescript(schema_path.read_text(encoding="utf-8"))
            ensure_calendar_schema(conn)
            ensure_calendar_revision_schema(conn)
            start_date, end_date = rows[0][0], rows[-1][0]
            conn.execute("DELETE FROM days WHERE date BETWEEN ? AND ?", (start_date, end_date))
            conn.executemany(
                "INSERT INTO days (date, kind, week_day) VALUES (?, ?, ?)",
                rows,
            )
            semester_ids = [
                row[0]
                for row in conn.execute(
                    """
                    SELECT id FROM semesters
                    WHERE start_date <= ? AND end_date >= ?
                    """,
                    (end_date, start_date),
                ).fetchall()
            ]
            for semester_id in semester_ids:
                conn.execute(
                    "INSERT OR IGNORE INTO calendar_revisions (semester_id) VALUES (?)",
                    (semester_id,),
                )
                conn.execute(
                    """
                    UPDATE calendar_revisions
                    SET revision = revision + 1, updated_at = datetime('now')
                    WHERE semester_id = ?
                    """,
                    (semester_id,),
                )
    finally:
        conn.close()

    counts = {}
    for _, kind, _ in rows:
        counts[kind] = counts.get(kind, 0) + 1
    return counts


def main(argv=None):
    parser = argparse.ArgumentParser(description="Import an academic calendar from an Excel workbook.")
    parser.add_argument("database", type=Path, help="SQLite database path")
    parser.add_argument("workbook", type=Path, help="Calendar.xlsx path")
    parser.add_argument("--schema", type=Path, default=BACKEND_DIR / "schema.sql")
    parser.add_argument("--academic-year-start", type=int)
    args = parser.parse_args(argv)

    try:
        counts = import_calendar(args.database, args.workbook, args.schema, args.academic_year_start)
    except (OSError, ValueError, sqlite3.Error) as exc:
        print(f"ERROR: {exc}")
        return 1
    summary = ", ".join(f"{kind}={count}" for kind, count in sorted(counts.items()))
    print(f"Imported academic calendar ({summary})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
