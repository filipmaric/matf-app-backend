#!/usr/bin/env python3
# Copyright (c) 2026 Filip Marić. See LICENCE.

"""Import course sessions from a workbook or exported CSV into SQLite."""

from __future__ import annotations

import argparse
import csv
import itertools
import re
import sqlite3
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from semester_utils import academic_year_label, parse_school_year_start_year, semester_display_name, school_year_semester_ids
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib.timetable_common import get_or_create_group, normalize_group_name
EXCLUDED_SHEETS = {"Фонд", "Све"}
CSV_TEACHER_USERNAME_FIELDS = ("teacher_username",)
CSV_TEACHER_NAME_FIELDS = ("teacher_name", "teacher")
CSV_COURSE_CODE_FIELDS = ("course_code", "course")
CSV_SEMESTER_FIELDS = ("semester_name", "semester_label", "semester", "term")
CSV_TYPE_FIELDS = ("type", "session_type")
CSV_GROUP_FIELDS = ("group_names", "groups")
CSV_WEEKLY_LESSONS_FIELDS = ("weekly_lessons", "weekly_lesson_count", "fond", "fund")
UNKNOWN_TEACHER_NAME = "Непознат наставник"
UNKNOWN_TEACHER_USERNAME = "unknown.teacher"


@dataclass(frozen=True)
class CourseSessionRow:
    sheet_name: str
    row_number: int
    teacher_name: str
    teacher_username: str
    course_code: str
    semester_label: str
    session_type: str
    group_names: tuple[str, ...]
    weekly_lessons: float


def normalize_text(value) -> str:
    return " ".join(str(value or "").split())


def parse_weekly_lessons(value) -> float:
    normalized = normalize_text(value).replace(",", ".")
    if not normalized:
        return 0.0
    try:
        weekly_lessons = float(normalized)
    except ValueError:
        return 0.0
    if weekly_lessons < 0:
        return 0.0
    return weekly_lessons


def row_has_content(row: tuple) -> bool:
    return any(normalize_text(value) for value in row)


def row_dict_has_content(row: dict[str, object]) -> bool:
    return any(normalize_text(value) for value in row.values())


def ensure_schema(conn: sqlite3.Connection, schema_path: Path) -> None:
    has_course_sessions = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'course_sessions'"
    ).fetchone()
    if has_course_sessions:
        return
    if not schema_path.exists():
        raise FileNotFoundError(f"Schema file not found: {schema_path}")
    conn.executescript(schema_path.read_text(encoding="utf-8"))


def record_skip(message: str, *, skipped_file) -> None:
    print(message, file=sys.stderr)
    if skipped_file is not None:
        skipped_file.write(message + "\n")


def clear_semester_schedule(conn: sqlite3.Connection, semester_id: int) -> None:
    cur = conn.cursor()
    cur.execute("PRAGMA foreign_keys = ON;")
    cur.execute(
        """
        DELETE FROM weekly_sessions
        WHERE session_id IN (
            SELECT id FROM course_sessions WHERE semester_id = ?
        )
        """,
        (semester_id,),
    )
    cur.execute(
        "DELETE FROM session_groups WHERE session_id IN (SELECT id FROM course_sessions WHERE semester_id = ?)",
        (semester_id,),
    )
    cur.execute(
        "DELETE FROM course_sessions WHERE semester_id = ?",
        (semester_id,),
    )


def split_group_names(value: str) -> tuple[str, ...]:
    names: list[str] = []
    for raw_part in re.split(r"[;,|]", normalize_text(value)):
        part = normalize_group_name(raw_part)
        if part and part not in names:
            names.append(part)
    return tuple(names)


def find_column(headers: tuple[str, ...], predicate) -> int | None:
    for index, header in enumerate(headers, start=1):
        if predicate(header):
            return index
    return None


def lookup_id(cur: sqlite3.Cursor, table: str, column: str, value: str) -> int | None:
    row = cur.execute(
        f"SELECT id FROM {table} WHERE trim({column}) = ?",
        (value,),
    ).fetchone()
    return int(row[0]) if row is not None else None


def lookup_teacher_id(
    cur: sqlite3.Cursor, teacher_username: str, teacher_name: str
) -> int | None:
    if teacher_username:
        teacher_id = lookup_id(cur, "teachers", "username", teacher_username)
        if teacher_id is not None:
            return teacher_id
    if not teacher_name:
        return None

    name_parts = normalize_text(teacher_name).split()
    name_candidates = [" ".join(name_parts)]
    if len(name_parts) >= 3:
        name_candidates.append(f"{name_parts[0]} {name_parts[-1]}")
        name_candidates.extend(" ".join(pair) for pair in itertools.combinations(name_parts, 2))
    for candidate in dict.fromkeys(name_candidates):
        teacher_id = lookup_id(cur, "teachers", "name", candidate)
        if teacher_id is not None:
            return teacher_id
    return None


def ensure_groups(cur: sqlite3.Cursor, group_names: set[str]) -> None:
    for group_name in sorted(group_names, key=normalize_group_name):
        get_or_create_group(cur, group_name)


def ensure_unknown_teacher(cur: sqlite3.Cursor) -> int:
    """Return the placeholder teacher used for unresolved workbook labels."""
    row = cur.execute(
        "SELECT id FROM teachers WHERE username = ?",
        (UNKNOWN_TEACHER_USERNAME,),
    ).fetchone()
    if row is not None:
        return int(row[0])

    cur.execute(
        """
        INSERT INTO teachers (name, username)
        VALUES (?, ?)
        """,
        (UNKNOWN_TEACHER_NAME, UNKNOWN_TEACHER_USERNAME),
    )
    return int(cur.lastrowid)


def lookup_course(cur: sqlite3.Cursor, course_code: str) -> tuple[int, int] | None:
    row = cur.execute(
        "SELECT id, semester FROM courses WHERE trim(code) = ?",
        (course_code,),
    ).fetchone()
    if row is None:
        return None
    return int(row[0]), int(row[1])


def resolve_semester_term(raw_value: str) -> str | None:
    value = normalize_text(raw_value).casefold()
    if not value:
        return None
    if any(token in value for token in ("јесењ", "jesenj", "fall", "autumn")):
        return "јесењи"
    if any(token in value for token in ("пролећ", "prolec", "spring")):
        return "пролећни"
    return None


def resolve_semester_name(cur: sqlite3.Cursor, academic_year: str, raw_value: str) -> str | None:
    value = normalize_text(raw_value)
    if not value:
        return None

    start_year = parse_school_year_start_year(academic_year)
    if start_year is None:
        return None

    semester_term = resolve_semester_term(value)
    if semester_term is None:
        return None
    return semester_display_name(start_year, semester_term)


def expected_course_semester(cur: sqlite3.Cursor, academic_year: str, semester_name: str) -> int | None:
    normalized = normalize_text(semester_name).casefold()
    if "јесењ" in normalized:
        fall_semester_id, _ = school_year_semester_ids(cur, academic_year)
        return fall_semester_id
    if "пролећ" in normalized:
        _, spring_semester_id = school_year_semester_ids(cur, academic_year)
        return spring_semester_id
    return None


def pick_csv_value(row: dict[str, object], *field_names: str) -> str:
    for field_name in field_names:
        value = normalize_text(row.get(field_name))
        if value:
            return value
    return ""


def load_workbook_rows(workbook_path: Path) -> list[CourseSessionRow]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    rows: list[CourseSessionRow] = []

    for sheet_name in workbook.sheetnames:
        if sheet_name in EXCLUDED_SHEETS:
            continue
        ws = workbook[sheet_name]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = tuple(normalize_text(value).casefold() for value in header_row)
        type_col = find_column(headers, lambda header: header == "тип")
        weekly_lessons_col = find_column(headers, lambda header: header == "фонд")
        groups_cols = [
            index
            for index, header in enumerate(headers, start=1)
            if "групе" in header
        ]

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row_has_content(row):
                continue
            teacher_name = normalize_text(row[0] if len(row) > 0 else None)
            course_code = normalize_text(row[1] if len(row) > 1 else None)
            semester_label = normalize_text(row[3] if len(row) > 3 else None)
            session_type = normalize_text(row[type_col - 1] if type_col and len(row) >= type_col else None)
            weekly_lessons = parse_weekly_lessons(
                row[weekly_lessons_col - 1]
                if weekly_lessons_col and len(row) >= weekly_lessons_col
                else None
            )
            group_names: list[str] = []
            for groups_col in groups_cols:
                if len(row) < groups_col:
                    continue
                for group_name in split_group_names(row[groups_col - 1]):
                    if group_name not in group_names:
                        group_names.append(group_name)

            rows.append(
                CourseSessionRow(
                    sheet_name=sheet_name,
                    row_number=row_number,
                    teacher_name=teacher_name,
                    teacher_username="",
                    course_code=course_code,
                    semester_label=semester_label,
                    session_type=session_type,
                    group_names=tuple(group_names),
                    weekly_lessons=weekly_lessons,
                )
            )

    return rows


def load_csv_rows(csv_path: Path) -> list[CourseSessionRow]:
    rows: list[CourseSessionRow] = []
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            return rows

        fieldnames = {normalize_text(name).casefold() for name in reader.fieldnames}
        aliases_present = {
            "teacher": any(name in fieldnames for name in CSV_TEACHER_NAME_FIELDS)
            or any(name in fieldnames for name in CSV_TEACHER_USERNAME_FIELDS),
            "course_code": any(name in fieldnames for name in CSV_COURSE_CODE_FIELDS),
            "semester_name": any(name in fieldnames for name in CSV_SEMESTER_FIELDS),
            "type": any(name in fieldnames for name in CSV_TYPE_FIELDS),
            "group_names": any(name in fieldnames for name in CSV_GROUP_FIELDS),
        }
        if not all(aliases_present.values()):
            missing = sorted(name for name, present in aliases_present.items() if not present)
            raise ValueError(
                "CSV input does not contain the expected course-session columns: "
                f"{', '.join(missing)}"
            )

        for row_number, row in enumerate(reader, start=2):
            if not row_dict_has_content(row):
                continue
            rows.append(
                CourseSessionRow(
                    sheet_name=csv_path.stem,
                    row_number=row_number,
                    teacher_name=pick_csv_value(row, *CSV_TEACHER_NAME_FIELDS),
                    teacher_username=pick_csv_value(row, *CSV_TEACHER_USERNAME_FIELDS),
                    course_code=pick_csv_value(row, *CSV_COURSE_CODE_FIELDS),
                    semester_label=pick_csv_value(row, *CSV_SEMESTER_FIELDS),
                    session_type=pick_csv_value(row, *CSV_TYPE_FIELDS),
                    group_names=split_group_names(pick_csv_value(row, *CSV_GROUP_FIELDS)),
                    weekly_lessons=parse_weekly_lessons(
                        pick_csv_value(row, *CSV_WEEKLY_LESSONS_FIELDS)
                    ),
                )
            )
    return rows


def load_rows(source_path: Path) -> list[CourseSessionRow]:
    suffix = source_path.suffix.lower()
    if suffix == ".csv":
        return load_csv_rows(source_path)
    return load_workbook_rows(source_path)


def plan_rows(
    cur: sqlite3.Cursor,
    academic_year: str,
    rows: list[CourseSessionRow],
    *,
    skipped_file=None,
) -> tuple[dict[int, list[CourseSessionRow]], dict[int, str], int, int]:
    planned_rows: dict[int, list[CourseSessionRow]] = {}
    semester_names_by_id: dict[int, str] = {}
    selected_rows = 0
    skipped_rows = 0

    for row in rows:
        problems: list[str] = []

        teacher_identifier = row.teacher_username or row.teacher_name
        if not teacher_identifier:
            problems.append("missing teacher")
        if not row.course_code:
            problems.append("missing course code")
        if not row.session_type:
            problems.append("missing type")
        if not row.group_names:
            problems.append("missing groups")

        semester_name = resolve_semester_name(cur, academic_year, row.semester_label)
        if semester_name is None:
            problems.append(f"unknown semester label: {row.semester_label or '<empty>'}")

        if problems:
            record_skip(
                f"SKIP\tsheet={row.sheet_name}\trow={row.row_number}\treason={'; '.join(problems)}",
                skipped_file=skipped_file,
            )
            skipped_rows += 1
            continue

        semester_term = resolve_semester_term(semester_name)
        start_year = parse_school_year_start_year(academic_year)
        if start_year is None or semester_term is None:
            semester_id = None
        else:
            row_match = cur.execute(
                """
                SELECT id
                FROM semesters
                WHERE academic_year_start = ? AND season = ?
                """,
                (start_year, semester_term),
            ).fetchone()
            semester_id = int(row_match[0]) if row_match is not None else None

        if semester_id is None:
            record_skip(
                f"SKIP\tsheet={row.sheet_name}\trow={row.row_number}\treason=unknown semester: {semester_name}",
                skipped_file=skipped_file,
            )
            skipped_rows += 1
            continue

        planned_rows.setdefault(semester_id, []).append(row)
        semester_names_by_id[semester_id] = semester_name
        selected_rows += 1

    return planned_rows, semester_names_by_id, selected_rows, skipped_rows


def import_rows(
    cur: sqlite3.Cursor,
    rows: list[CourseSessionRow],
    semester_id: int,
    semester_name: str,
    academic_year: str,
    *,
    skipped_file=None,
) -> tuple[int, int, int]:
    imported_sessions = 0
    imported_group_links = 0
    skipped_rows = 0
    seen_keys: set[tuple[int, int, int, str, tuple[int, ...]]] = set()
    target_course_semester = expected_course_semester(cur, academic_year, semester_name)

    for row in rows:
        problems: list[str] = []
        teacher_identifier = row.teacher_username or row.teacher_name

        if not teacher_identifier:
            problems.append("missing teacher")
        if not row.course_code:
            problems.append("missing course code")
        if not row.session_type:
            problems.append("missing type")
        if not row.group_names:
            problems.append("missing groups")

        if problems:
            record_skip(
                f"SKIP\tsheet={row.sheet_name}\trow={row.row_number}\treason={'; '.join(problems)}",
                skipped_file=skipped_file,
            )
            skipped_rows += 1
            continue

        course_row = lookup_course(cur, row.course_code)
        if course_row is None:
            problems.append(f"unknown course: {row.course_code}")
            course_id = None
        else:
            course_id, course_semester = course_row
            if target_course_semester is None:
                problems.append(f"unsupported semester: {semester_name}")
            elif course_semester != target_course_semester:
                problems.append(
                    f"semester mismatch for course {row.course_code}: "
                    f"course semester={course_semester}, row semester={target_course_semester}"
                )

        teacher_id = None
        teacher_id = lookup_teacher_id(cur, row.teacher_username, row.teacher_name)
        if teacher_id is None:
            teacher_id = ensure_unknown_teacher(cur)

        group_ids: list[int] = []
        missing_groups: list[str] = []
        for group_name in row.group_names:
            group_id = lookup_id(cur, "groups", "name", group_name)
            if group_id is None:
                get_or_create_group(cur, group_name)
                group_id = lookup_id(cur, "groups", "name", group_name)
            if group_id is None:
                missing_groups.append(group_name)
                continue
            if group_id not in group_ids:
                group_ids.append(group_id)
        if missing_groups:
            problems.append(f"unknown groups: {', '.join(missing_groups)}")

        if problems:
            print(
                f"SKIP\tsheet={row.sheet_name}\trow={row.row_number}\treason={'; '.join(problems)}",
                file=sys.stderr,
            )
            skipped_rows += 1
            continue

        session_key = (
            course_id,
            teacher_id,
            semester_id,
            row.session_type,
            tuple(group_ids),
        )
        if session_key in seen_keys:
            print(
                f"SKIP\tsheet={row.sheet_name}\trow={row.row_number}\treason=duplicate session",
                file=sys.stderr,
            )
            skipped_rows += 1
            continue
        seen_keys.add(session_key)

        cur.execute(
            """
            INSERT INTO course_sessions
                (course_id, teacher_id, semester_id, type, weekly_lessons)
            VALUES (?, ?, ?, ?, ?)
            """,
            (course_id, teacher_id, semester_id, row.session_type, row.weekly_lessons),
        )
        session_id = int(cur.lastrowid)
        imported_sessions += 1

        for group_id in group_ids:
            cur.execute(
                """
                INSERT INTO session_groups (session_id, group_id)
                VALUES (?, ?)
                """,
                (session_id, group_id),
            )
            imported_group_links += 1

    return imported_sessions, imported_group_links, skipped_rows


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="Import course sessions from a workbook or exported CSV into a SQLite database."
    )
    parser.add_argument(
        "--year",
        required=True,
        help="School year start year like 2026",
    )
    parser.add_argument(
        "--source",
        "--workbook",
        dest="source",
        required=True,
        help="Path to the podela workbook (.xlsx) or exported CSV",
    )
    parser.add_argument(
        "--database",
        required=True,
        help="SQLite database path",
    )
    parser.add_argument(
        "--schema",
        required=True,
        help="Path to backend/schema.sql used when the database is empty",
    )
    parser.add_argument(
        "--failed-rows-file",
        default=None,
        help="Write skipped rows to this file in addition to stderr",
    )
    args = parser.parse_args(argv)

    source_path = Path(args.source)
    database_path = Path(args.database)
    schema_path = Path(args.schema)

    if not source_path.exists():
        print(f"Source file not found: {source_path}", file=sys.stderr)
        return 2

    start_year = parse_school_year_start_year(args.year)
    if start_year is None:
        print(f"Invalid year: {args.year!r}", file=sys.stderr)
        return 2

    rows = load_rows(source_path)
    if not rows:
        print("No course session rows found.", file=sys.stderr)
        return 1

    database_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(database_path)
    skipped_file_handle = None
    try:
        if args.failed_rows_file:
            failed_rows_path = Path(args.failed_rows_file)
            failed_rows_path.parent.mkdir(parents=True, exist_ok=True)
            skipped_file_handle = failed_rows_path.open("w", encoding="utf-8")
        conn.execute("PRAGMA foreign_keys = ON;")
        ensure_schema(conn, schema_path)
        with conn:
            cur = conn.cursor()
            planned_rows, semester_names_by_id, selected_rows, skipped_rows = plan_rows(
                cur,
                start_year,
                rows,
                skipped_file=skipped_file_handle,
            )
            all_group_names = {
                group_name
                for semester_rows in planned_rows.values()
                for row in semester_rows
                for group_name in row.group_names
            }
            ensure_groups(cur, all_group_names)
            for semester_id in sorted(planned_rows):
                clear_semester_schedule(conn, semester_id)

            imported_sessions = 0
            imported_group_links = 0
            semester_summaries: list[tuple[str, int, int, int]] = []
            for semester_id in sorted(planned_rows):
                semester_rows = planned_rows[semester_id]
                semester_name = semester_names_by_id[semester_id]
                semester_imported_sessions, semester_imported_group_links, semester_skipped = (
                    import_rows(
                        cur,
                    semester_rows,
                    semester_id,
                    semester_name,
                    start_year,
                    skipped_file=skipped_file_handle,
                )
                )
                imported_sessions += semester_imported_sessions
                imported_group_links += semester_imported_group_links
                skipped_rows += semester_skipped
                semester_summaries.append(
                    (
                        semester_name,
                        semester_imported_sessions,
                        semester_imported_group_links,
                        semester_skipped,
                    )
                )
    except Exception as exc:
        conn.rollback()
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        if skipped_file_handle is not None:
            skipped_file_handle.close()
        conn.close()

    print(f"YEAR={start_year}")
    print(f"ACADEMIC_YEAR={academic_year_label(start_year)}")
    print(f"ROWS={len(rows)}")
    print(f"SELECTED={selected_rows}")
    for semester_name, semester_imported_sessions, semester_imported_group_links, semester_skipped in semester_summaries:
        print(f"SEMESTER={semester_name}")
        print(f"IMPORTED={semester_imported_sessions}")
        print(f"SESSION_GROUPS={semester_imported_group_links}")
        print(f"SKIPPED={semester_skipped}")
    print(f"TOTAL_IMPORTED={imported_sessions}")
    print(f"TOTAL_SESSION_GROUPS={imported_group_links}")
    print(f"TOTAL_SKIPPED={skipped_rows}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
