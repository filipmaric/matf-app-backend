# Copyright (c) 2026 Filip Marić. See LICENCE.
import sqlite3
import subprocess
import sys
import importlib
from pathlib import Path

from openpyxl import Workbook, load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
PYTHON = sys.executable
sys.path.insert(0, str(REPO_ROOT / "scripts"))
from lib.timetable_common import normalize_group_name

load_workbook_rows = importlib.import_module(
    "scripts.import.import_course_sessions"
).load_workbook_rows
lookup_teacher_id = importlib.import_module(
    "scripts.import.import_course_sessions"
).lookup_teacher_id


def semester_number_for_token(semester_token):
    if semester_token == "јесењи":
        return 3
    if semester_token == "пролећни":
        return 4
    raise ValueError(f"Unknown semester token: {semester_token}")


def write_course_sessions_workbook(path):
    workbook = Workbook()
    workbook.remove(workbook.active)

    fall = workbook.create_sheet("КРИ")
    fall.append(["Наставник", "Шифра", "Назив", "Семестар", "Тип", "Фонд", "Групе"])
    fall.append([])
    fall.append(["Fall Teacher", "FALL1", "Fall Course", "јесењи", "п", 2, "3af, 4ai"])
    fall.append(["Other Teacher", "FALL2", "Other Fall Course", "јесењи", "в", 1, "3bf"])
    fall.append(["Incomplete Teacher", "INCOMPLETE", "Incomplete Course", "јесењи", None, 1, "3cf"])

    spring = workbook.create_sheet("ПРО")
    spring.append(["Наставник", "Шифра", "Назив", "Семестар", "Тип", "Фонд", "Групе"])
    spring.append(["Spring Teacher", "SPRING1", "Spring Course", "пролећни", "в", 2, "1o1"])
    spring.append(["Another Spring Teacher", "SPRING2", "Another Spring Course", "пролећни", "п", 1, "1o2"])

    workbook.save(path)


def test_workbook_import_merges_regular_and_empty_groups(tmp_path):
    workbook_path = tmp_path / "course_sessions.xlsx"
    write_course_sessions_workbook(workbook_path)
    rows = load_workbook_rows(workbook_path)

    row = next(row for row in rows if row.sheet_name == "КРИ" and row.row_number == 3)

    assert "3af" in row.group_names
    assert "4ai" in row.group_names
    assert len(row.group_names) == len(set(row.group_names))


def test_three_part_teacher_name_falls_back_to_two_name_match():
    conn = sqlite3.connect(":memory:")
    conn.execute("CREATE TABLE teachers (id INTEGER PRIMARY KEY, name TEXT, username TEXT)")
    conn.execute(
        "INSERT INTO teachers (name, username) VALUES (?, ?)",
        ("Милена Јаничић", "milena.janicic"),
    )
    try:
        teacher_id = lookup_teacher_id(conn.cursor(), "", "Милена Вујошевић Јаничић")
    finally:
        conn.close()

    assert teacher_id == 1


def init_db(db_path):
    conn = sqlite3.connect(db_path)
    try:
        with conn:
            with open(REPO_ROOT / "schema.sql", encoding="utf-8") as f:
                conn.executescript(f.read())
            conn.execute(
                """
                INSERT INTO semesters (id, academic_year_start, season, start_date, end_date)
                VALUES (?, ?, ?, ?, ?)
                """,
                (2, 2025, "пролећни", "2026-03-23", "2026-09-30"),
            )
            conn.execute(
                """
                INSERT INTO semesters (id, academic_year_start, season, start_date, end_date)
                VALUES (?, ?, ?, ?, ?)
                """,
                (3, 2026, "јесењи", "2026-10-01", "2027-02-28"),
            )
            conn.execute(
                """
                INSERT INTO semesters (id, academic_year_start, season, start_date, end_date)
                VALUES (?, ?, ?, ?, ?)
                """,
                (4, 2026, "пролећни", "2027-03-01", "2027-09-30"),
            )
            conn.execute(
                """
                INSERT INTO teachers (name, username)
                VALUES (?, ?)
                """,
                ("Old Teacher", "old.teacher"),
            )
            conn.execute(
                """
                INSERT INTO courses (name, code, semester)
                VALUES (?, ?, ?)
                """,
                ("Old Course", "OLD1", 2),
            )
            conn.execute(
                """
                INSERT INTO course_sessions (course_id, teacher_id, semester_id, type)
                VALUES (
                    (SELECT id FROM courses WHERE code = ?),
                    (SELECT id FROM teachers WHERE username = ?),
                    (SELECT id FROM semesters WHERE academic_year_start = ? AND season = ?),
                    ?
                )
                """,
                ("OLD1", "old.teacher", 2025, "пролећни", "п"),
            )
            conn.execute(
                """
                INSERT INTO session_groups (session_id, group_id)
                VALUES (
                    (SELECT cs.id FROM course_sessions cs
                     JOIN courses c ON c.id = cs.course_id
                     WHERE c.code = ? AND cs.type = ?),
                    (SELECT id FROM groups WHERE name = ?)
                )
                """,
                ("OLD1", "п", "3s15"),
            )
    finally:
        conn.close()


def collect_matching_target_rows(workbook_path, semester_token):
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    samples = []
    for sheet_name in workbook.sheetnames:
        if sheet_name in {"Фонд", "Све"}:
            continue
        ws = workbook[sheet_name]
        headers = [str(cell.value or "").strip().casefold() for cell in ws[1]]
        type_col = next((i for i, header in enumerate(headers) if header == "тип"), None)
        groups_col = next(
            (i for i, header in enumerate(headers) if "групе" in header and "празне" not in header),
            None,
        )
        semester_col = next(
            (i for i, header in enumerate(headers) if "семестар" in header or header == "једној"),
            None,
        )
        for row in ws.iter_rows(min_row=2, values_only=True):
            semester_label = (
                str(row[semester_col] or "").strip()
                if semester_col is not None and len(row) > semester_col
                else ""
            )
            if semester_token not in semester_label:
                continue
            teacher_name = str(row[0] or "").strip() if len(row) > 0 else ""
            course_code = str(row[1] or "").strip() if len(row) > 1 else ""
            course_name = str(row[2] or "").strip() if len(row) > 2 else ""
            session_type = (
                str(row[type_col] or "").strip()
                if type_col is not None and len(row) > type_col
                else ""
            )
            groups_text = (
                str(row[groups_col] or "").strip()
                if groups_col is not None and len(row) > groups_col
                else ""
            )
            if teacher_name and course_code and course_name and session_type and groups_text:
                samples.append((teacher_name, course_code, course_name, session_type, groups_text))

    assert samples, "No importable target semester row found in workbook"
    return samples


def seed_matching_target_row(db_path, sample, course_semester):
    teacher_name, course_code, course_name, session_type, groups_text = sample
    conn = sqlite3.connect(db_path)
    try:
        with conn:
            conn.execute(
                """
                INSERT OR IGNORE INTO teachers (name, username)
                VALUES (?, ?)
                """,
                (teacher_name, "target.teacher"),
            )
            conn.execute(
                """
                INSERT OR IGNORE INTO courses (name, code, semester)
                VALUES (?, ?, ?)
                """,
                (course_name, course_code, course_semester),
            )
            for group_name in {
                normalize_group_name(part)
                for part in groups_text.split(",")
                if normalize_group_name(part)
            }:
                conn.execute(
                    """
                    INSERT OR IGNORE INTO groups (name, description)
                    VALUES (?, NULL)
                    """,
                    (group_name,),
                )
    finally:
        conn.close()


def run_script(args):
    return subprocess.run(
        [PYTHON, *args],
        cwd=REPO_ROOT,
        text=True,
        capture_output=True,
        check=True,
    )


def test_import_course_sessions_from_podela_imports_rows_and_reports_missing(tmp_path):
    db_path = tmp_path / "course_sessions.db"
    skipped_path = tmp_path / "skipped_rows.log"
    workbook_path = tmp_path / "course_sessions.xlsx"
    write_course_sessions_workbook(workbook_path)
    init_db(db_path)
    fall_samples = collect_matching_target_rows(workbook_path, "јесењи")
    spring_samples = collect_matching_target_rows(workbook_path, "пролећни")
    seed_matching_target_row(db_path, fall_samples[0], semester_number_for_token("јесењи"))
    seed_matching_target_row(db_path, spring_samples[0], semester_number_for_token("пролећни"))
    seeded_codes = {fall_samples[0][1], spring_samples[0][1]}
    mismatch_sample = next(
        sample
        for sample in spring_samples[1:] + fall_samples[1:]
        if sample[1] not in seeded_codes
    )
    seed_matching_target_row(
        db_path,
        mismatch_sample,
        semester_number_for_token("јесењи"),
    )

    result = run_script(
        [
            "scripts/import/import_course_sessions.py",
            "--year",
            "2026",
            "--database",
            str(db_path),
            "--workbook",
            str(workbook_path),
            "--schema",
            str(REPO_ROOT / "schema.sql"),
            "--failed-rows-file",
            str(skipped_path),
        ]
    )

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        prior_rows = conn.execute(
            """
            SELECT c.code
            FROM course_sessions cs
            JOIN courses c ON c.id = cs.course_id
            JOIN semesters s ON s.id = cs.semester_id
            WHERE s.academic_year_start = ? AND s.season = ?
            ORDER BY c.code
            """,
            (2025, "пролећни"),
        ).fetchall()
        target_fall_rows = conn.execute(
            """
            SELECT c.code
            FROM course_sessions cs
            JOIN courses c ON c.id = cs.course_id
            JOIN semesters s ON s.id = cs.semester_id
            WHERE s.academic_year_start = ? AND s.season = ?
            ORDER BY c.code
            """,
            (2026, "јесењи"),
        ).fetchall()
        target_spring_rows = conn.execute(
            """
            SELECT c.code
            FROM course_sessions cs
            JOIN courses c ON c.id = cs.course_id
            JOIN semesters s ON s.id = cs.semester_id
            WHERE s.academic_year_start = ? AND s.season = ?
            ORDER BY c.code
            """,
            (2026, "пролећни"),
        ).fetchall()
    finally:
        conn.close()

    assert [row["code"] for row in prior_rows] == ["OLD1"]
    assert "OLD1" not in {row["code"] for row in target_fall_rows}
    assert "OLD1" not in {row["code"] for row in target_spring_rows}
    assert target_fall_rows
    assert target_spring_rows
    assert "SKIP\t" in result.stderr
    assert "YEAR=2026" in result.stdout
    assert "ACADEMIC_YEAR=2026/27" in result.stdout
    assert "SEMESTER=2026/27. јесењи" in result.stdout
    assert "SEMESTER=2026/27. пролећни" in result.stdout
    assert "semester mismatch for course" in result.stderr
    assert "TOTAL_IMPORTED=" in result.stdout
    assert "TOTAL_SESSION_GROUPS=" in result.stdout
    assert "ROWS=" in result.stdout
    assert "SELECTED=" in result.stdout
    assert skipped_path.exists()
    assert "SKIP\t" in skipped_path.read_text(encoding="utf-8")


def test_import_course_sessions_uses_placeholder_for_unknown_teacher(tmp_path):
    db_path = tmp_path / "course_sessions.db"
    csv_path = tmp_path / "course_sessions.csv"
    init_db(db_path)
    csv_path.write_text(
        "\n".join(
            [
                "teacher_name,course_code,semester_name,type,group_names",
                "Missing Teacher,NM1,јесењи,в,1o4",
            ]
        ),
        encoding="utf-8",
    )

    conn = sqlite3.connect(db_path)
    try:
        with conn:
            conn.execute(
                """
                INSERT INTO courses (name, code, semester)
                VALUES (?, ?, ?)
                """,
                ("Known Course", "NM1", semester_number_for_token("јесењи")),
            )
    finally:
        conn.close()

    result = run_script(
        [
            "scripts/import/import_course_sessions.py",
            "--year",
            "2026",
            "--database",
            str(db_path),
            "--source",
            str(csv_path),
            "--schema",
            str(REPO_ROOT / "schema.sql"),
        ]
    )

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        row = conn.execute(
            """
            SELECT c.code, t.name, t.username, g.name AS group_name
            FROM course_sessions cs
            JOIN courses c ON c.id = cs.course_id
            JOIN teachers t ON t.id = cs.teacher_id
            JOIN session_groups sg ON sg.session_id = cs.id
            JOIN groups g ON g.id = sg.group_id
            WHERE c.code = ?
            """,
            ("NM1",),
        ).fetchone()
    finally:
        conn.close()

    assert "TOTAL_IMPORTED=1" in result.stdout
    assert "unknown teacher" not in result.stderr
    assert row["name"] == "Непознат наставник"
    assert row["username"] == "unknown.teacher"
    assert row["group_name"] == "1o4"
