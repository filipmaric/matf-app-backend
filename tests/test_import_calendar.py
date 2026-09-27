# Copyright (c) 2026 Filip Marić. See LICENCE.
from datetime import date, timedelta
import importlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

import_calendar = importlib.import_module("scripts.import.import_calendar").import_calendar


ROOT = Path(__file__).resolve().parents[1]


def test_import_calendar_reads_kinds_and_makeup_comments(tmp_path):
    workbook_path = tmp_path / "Calendar.xlsx"
    database_path = tmp_path / "calendar.db"

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Календар за школску 2026/27. годину"
    start = date(2026, 10, 1)
    fills = {
        "teaching": PatternFill("solid", fgColor="FFFFFF00"),
        "non_working": PatternFill("solid", fgColor="FFFF0000"),
        "exam": PatternFill("solid", fgColor="FF4285F4"),
        "colloquium": PatternFill("solid", fgColor="FFFFF2CC"),
        "makeup": PatternFill("solid", fgColor="FFFF9900"),
    }
    kinds = {
        date(2026, 10, 1): "teaching",
        date(2026, 10, 2): "exam",
        date(2026, 10, 3): "non_working",
        date(2026, 10, 4): "colloquium",
        date(2026, 10, 5): "makeup",
    }
    current = start
    last = date(2027, 9, 30)
    while current <= last:
        offset = (current - start).days
        position = offset + 3  # the first week starts in column F (Thursday)
        row = 3 + position // 7
        column = 3 + position % 7
        cell = sheet.cell(row, column, current.day)
        kind = kinds.get(current, "non_working")
        cell.fill = fills[kind]
        if kind == "makeup":
            cell.comment = Comment("Петак", "test")
        current += timedelta(days=1)
    workbook.save(workbook_path)

    counts = import_calendar(database_path, workbook_path, ROOT / "schema.sql")

    assert counts == {
        "colloquium": 1,
        "exam": 1,
        "makeup": 1,
        "non_working": 361,
        "teaching": 1,
    }
    import sqlite3

    conn = sqlite3.connect(database_path)
    rows = conn.execute(
        "SELECT date, kind, week_day FROM days WHERE kind IN ('exam', 'colloquium', 'makeup') ORDER BY date"
    ).fetchall()
    conn.close()
    assert rows == [
        ("2026-10-02", "exam", -1),
        ("2026-10-04", "colloquium", -1),
        ("2026-10-05", "makeup", 4),
    ]
